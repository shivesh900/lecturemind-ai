"""Exports: TXT / SRT / Markdown / DOCX / PDF / CSV / JSON, built from the real
stored content."""
from __future__ import annotations

import csv
import io
import json
import os
import re
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

from ..utils.text import format_ts

# ---------- fonts (Unicode-capable when DejaVu is installed) ----------
_FONT, _FONT_BOLD = "Helvetica", "Helvetica-Bold"
for _dir in ("/usr/share/fonts/truetype/dejavu", "/usr/share/fonts/dejavu", "/Library/Fonts", "C:/Windows/Fonts"):
    reg, bold = os.path.join(_dir, "DejaVuSans.ttf"), os.path.join(_dir, "DejaVuSans-Bold.ttf")
    if os.path.exists(reg) and os.path.exists(bold):
        try:
            pdfmetrics.registerFont(TTFont("DejaVu", reg))
            pdfmetrics.registerFont(TTFont("DejaVu-Bold", bold))
            _FONT, _FONT_BOLD = "DejaVu", "DejaVu-Bold"
        except Exception:
            pass
        break


def _pdf_text(s: str) -> str:
    s = s or ""
    if _FONT == "Helvetica":
        s = s.encode("latin-1", "replace").decode("latin-1")
    return s


# ---------- transcript ----------
def transcript_txt(lecture: dict, segments: list[dict], timestamps: bool = True) -> str:
    head = [lecture.get("title", "Lecture"), f"Duration: {format_ts(lecture.get('duration'))}", ""]
    lines = []
    for s in segments:
        prefix = f"[{format_ts(s['start'])}] " if timestamps else ""
        spk = f"{s['speaker']}: " if s.get("speaker") else ""
        lines.append(f"{prefix}{spk}{s['text']}")
    return "\n".join(head + lines) + "\n"


def transcript_srt(segments: list[dict]) -> str:
    def t(x: float) -> str:
        ms = int(round(x * 1000))
        h, ms = divmod(ms, 3600000)
        m, ms = divmod(ms, 60000)
        s, ms = divmod(ms, 1000)
        return f"{h:02d}:{m:02d}:{s:02d},{ms:03d}"
    blocks = [f"{i + 1}\n{t(s['start'])} --> {t(s['end'])}\n{s['text']}\n" for i, s in enumerate(segments)]
    return "\n".join(blocks)


def transcript_docx(lecture: dict, segments: list[dict]) -> bytes:
    doc = Document()
    _docx_base(doc)
    doc.add_heading(lecture.get("title", "Lecture"), 0)
    meta = doc.add_paragraph()
    meta.add_run(f"Duration {format_ts(lecture.get('duration'))} · language {lecture.get('language') or 'auto'}").italic = True
    for s in segments:
        p = doc.add_paragraph()
        r = p.add_run(f"[{format_ts(s['start'])}] ")
        r.bold = True
        r.font.color.rgb = RGBColor(0x4F, 0x46, 0xE5)
        if s.get("speaker"):
            p.add_run(f"{s['speaker']}: ").bold = True
        p.add_run(s["text"])
    return _docx_bytes(doc)


# ---------- markdown -> docx / pdf ----------
_INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")


def _docx_base(doc: Document) -> None:
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)


def _docx_bytes(doc: Document) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _docx_inline(p, text: str) -> None:
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            p.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = p.add_run(part[1:-1])
            r.font.name = "Consolas"
        else:
            p.add_run(part)


def markdown_to_docx(md: str, title: str | None = None) -> bytes:
    doc = Document()
    _docx_base(doc)
    if title and not md.lstrip().startswith("# "):
        doc.add_heading(title, 0)
    for raw in md.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("# "):
            doc.add_heading(line[2:].replace("`", ""), 0)
        elif line.startswith("## "):
            doc.add_heading(line[3:].replace("`", ""), 1)
        elif line.startswith("### "):
            doc.add_heading(line[4:].replace("`", ""), 2)
        elif re.match(r"^\s*[-*] ", line):
            _docx_inline(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*] ", "", line))
        elif re.match(r"^\s*\d+\. ", line):
            _docx_inline(doc.add_paragraph(style="List Number"), re.sub(r"^\s*\d+\. ", "", line))
        elif line.startswith("> "):
            p = doc.add_paragraph()
            r = p.add_run(line[2:])
            r.italic = True
        else:
            _docx_inline(doc.add_paragraph(), line)
    return _docx_bytes(doc)


def _styles():
    ss = getSampleStyleSheet()
    base = ParagraphStyle("base", parent=ss["Normal"], fontName=_FONT, fontSize=10.5, leading=15, alignment=TA_LEFT)
    return {
        "title": ParagraphStyle("t", parent=base, fontName=_FONT_BOLD, fontSize=20, leading=25, spaceAfter=10,
                                textColor=colors.HexColor("#1e1b4b")),
        "h1": ParagraphStyle("h1", parent=base, fontName=_FONT_BOLD, fontSize=15, leading=20, spaceBefore=12,
                             spaceAfter=6, textColor=colors.HexColor("#3730a3")),
        "h2": ParagraphStyle("h2", parent=base, fontName=_FONT_BOLD, fontSize=12.5, leading=17, spaceBefore=8,
                             spaceAfter=4, textColor=colors.HexColor("#4338ca")),
        "body": base,
        "quote": ParagraphStyle("q", parent=base, textColor=colors.HexColor("#6b7280"), leftIndent=8),
        "small": ParagraphStyle("s", parent=base, fontSize=9, textColor=colors.HexColor("#6b7280")),
    }


def _pdf_inline(text: str) -> str:
    out = []
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            out.append(f"<font name='{_FONT_BOLD}'>{escape(_pdf_text(part[2:-2]))}</font>")
        elif part.startswith("`") and part.endswith("`"):
            out.append(f"<font name='Courier' color='#4338ca'>{escape(_pdf_text(part[1:-1]))}</font>")
        else:
            out.append(escape(_pdf_text(part)))
    return "".join(out)


def markdown_to_pdf(md: str, title: str | None = None, footer: str = "Generated with LectureMind AI") -> bytes:
    st = _styles()
    story: list = []
    if title and not md.lstrip().startswith("# "):
        story.append(Paragraph(escape(_pdf_text(title)), st["title"]))
    bullets: list = []
    numbered: list = []

    def flush():
        nonlocal bullets, numbered
        if bullets:
            story.append(ListFlowable([ListItem(b, leftIndent=12) for b in bullets], bulletType="bullet", start="•",
                                      leftIndent=12, bulletFontName=_FONT))
            bullets = []
        if numbered:
            story.append(ListFlowable([ListItem(b, leftIndent=14) for b in numbered], bulletType="1", leftIndent=14,
                                      bulletFontName=_FONT))
            numbered = []

    for raw in md.splitlines():
        line = raw.rstrip()
        if re.match(r"^\s*[-*] ", line):
            if numbered:
                flush()
            bullets.append(Paragraph(_pdf_inline(re.sub(r"^\s*[-*] ", "", line)), st["body"]))
            continue
        if re.match(r"^\s*\d+\. ", line):
            if bullets:
                flush()
            numbered.append(Paragraph(_pdf_inline(re.sub(r"^\s*\d+\. ", "", line)), st["body"]))
            continue
        flush()
        if not line.strip():
            story.append(Spacer(1, 3))
        elif line.startswith("# "):
            story.append(Paragraph(_pdf_inline(line[2:]), st["title"]))
        elif line.startswith("## "):
            story.append(Paragraph(_pdf_inline(line[3:]), st["h1"]))
        elif line.startswith("### "):
            story.append(Paragraph(_pdf_inline(line[4:]), st["h2"]))
        elif line.startswith("> "):
            story.append(Paragraph(_pdf_inline(line[2:]), st["quote"]))
        else:
            story.append(Paragraph(_pdf_inline(line), st["body"]))
    flush()

    buf = io.BytesIO()

    def _footer(canvas, doc):
        canvas.saveState()
        canvas.setFont(_FONT, 8)
        canvas.setFillColor(colors.HexColor("#9ca3af"))
        canvas.drawString(18 * mm, 10 * mm, footer)
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, f"Page {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=18 * mm,
                            bottomMargin=18 * mm, title=title or "LectureMind export")
    doc.build(story or [Paragraph("(empty)", st["body"])], onFirstPage=_footer, onLaterPages=_footer)
    return buf.getvalue()


# ---------- summary / quiz / flashcards ----------
def summary_markdown(lecture: dict, summary_doc: dict) -> str:
    s = summary_doc.get("summary", {})
    md = [f"# {lecture.get('title', 'Lecture')} — Summary"]
    if summary_doc.get("mode") == "demo":
        md.append("> Demo mode: extractive summary built without an AI model.")
    md += ["", "## Executive summary", s.get("executive_summary", "")]
    if s.get("quick_summary"):
        md += ["", "## Quick summary (5 lines)"] + [f"- {x}" for x in s["quick_summary"]]
    if summary_doc.get("chapters"):
        md += ["", "## Chapters"] + [f"- `{format_ts(c['start'])}` **{c['title']}** — {c.get('summary', '')[:300]}"
                                    for c in summary_doc["chapters"]]
    if s.get("detailed_summary"):
        md += ["", "## Detailed summary"]
        for sec in s["detailed_summary"]:
            md += ["", f"### {sec['heading']}", sec["content"]]
    for key, label in (("main_topics", "Main topics"), ("arguments", "Important arguments & explanations"),
                       ("key_takeaways", "Key takeaways"), ("examples", "Examples mentioned")):
        items = s.get(key) or []
        if items:
            md += ["", f"## {label}"]
            for it in items:
                md.append(f"- **{it['title']}**" + (f": {it['description']}" if it.get("description") else "")
                          if isinstance(it, dict) else f"- {it}")
    if s.get("definitions"):
        md += ["", "## Definitions"] + [f"- **{d['term']}**: {d['definition']}" for d in s["definitions"]]
    if s.get("formulas"):
        md += ["", "## Formulas & concepts"] + [f"- `{f['expression']}`" + (f" — {f['meaning']}" if f.get("meaning") else "")
                                               for f in s["formulas"]]
    return "\n".join(md).strip() + "\n"


def quiz_markdown(lecture: dict, quiz: dict, with_answers: bool = True) -> str:
    md = [f"# {quiz.get('title') or 'Quiz'}", f"Lecture: **{lecture.get('title', '')}** · difficulty: {quiz.get('difficulty', 'mixed')}", ""]
    labels = {"mcq": "Multiple choice", "true_false": "True / False", "short": "Short answer", "long": "Long answer"}
    n = 0
    for qtype in ("mcq", "true_false", "short", "long"):
        qs = [q for q in quiz.get("questions", []) if q["type"] == qtype]
        if not qs:
            continue
        md += [f"## {labels[qtype]}"]
        for q in qs:
            n += 1
            star = " ★" if q.get("important") else ""
            md.append(f"{n}. **{q['question']}** ({q.get('difficulty', 'medium')}){star}")
            for i, o in enumerate(q.get("options") or []):
                md.append(f"- {'ABCD'[i] if i < 4 else i + 1}) {o}")
            md.append("")
    if with_answers:
        md += ["## Answer key"]
        n = 0
        for qtype in ("mcq", "true_false", "short", "long"):
            for q in [q for q in quiz.get("questions", []) if q["type"] == qtype]:
                n += 1
                md.append(f"{n}. **{q['correct_answer']}** — {q.get('explanation', '')}")
    return "\n".join(md).strip() + "\n"


def flashcards_csv(cards: list[dict]) -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["front", "back", "difficulty", "your_rating", "times_reviewed"])
    for c in cards:
        w.writerow([c.get("front", ""), c.get("back", ""), c.get("difficulty", ""), c.get("rating") or "", c.get("review_count", 0)])
    return buf.getvalue()


def to_json_bytes(data: dict) -> bytes:
    return json.dumps(data, ensure_ascii=False, indent=2, default=str).encode("utf-8")
